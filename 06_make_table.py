# -*- coding: utf-8 -*-
"""Generate the Q1-paper-style comparison table (png/md/csv/xlsx) + figures.
Reads eval_results/combined_metrics.json + benchmark.json.
Entity colors follow the validated categorical palette, fixed per model."""
import json
import re
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import openpyxl
from openpyxl.styles import Font as XLFont, Alignment

EXP = Path(r"D:/pig-objection detection/experiments")
EVAL_DIR = EXP / "eval_results"
OUT_DIR = EXP / "outputs"
OUT_DIR.mkdir(exist_ok=True)

MODELS = ["frcnn", "ssd", "yolov5s", "yolov8s", "yolo11s", "rtdetr-l"]
DISPLAY = {
    "frcnn": ("Faster R-CNN", "ResNet-50-FPN"),
    "ssd": ("SSD300", "VGG-16"),
    "yolov5s": ("YOLOv5s", "CSPDarknet53"),
    "yolov8s": ("YOLOv8s", "C2f-Darknet"),
    "yolo11s": ("YOLOv11s", "C3k2-Darknet"),
    "rtdetr-l": ("RT-DETR-L", "HGNetv2"),
}
# validated categorical palette (fixed order, entity-fixed)
COLOR = {
    "frcnn": "#2a78d6",
    "ssd": "#eb6834",
    "yolov5s": "#1baf7a",
    "yolov8s": "#eda100",
    "yolo11s": "#e87ba4",
    "rtdetr-l": "#008300",
}
INK = "#0b0b0b"; MUTED = "#898781"; SURFACE = "#fcfcfb"; GRID = "#e1e0d9"

plt.rcParams.update({
    "font.family": "serif",
    "font.serif": ["Times New Roman", "DejaVu Serif"],
    "axes.edgecolor": "#c3c2b7",
    "axes.labelcolor": INK,
    "xtick.color": MUTED, "ytick.color": MUTED,
    "text.color": INK,
    "figure.facecolor": SURFACE, "axes.facecolor": SURFACE,
    "savefig.facecolor": SURFACE,
    "savefig.dpi": 300,
})

metrics = json.loads((EVAL_DIR / "combined_metrics.json").read_text())
bench = json.loads((EVAL_DIR / "benchmark.json").read_text())


def train_time(name):
    if name in ("ssd", "frcnn"):
        log = EXP / "eval_results" / f"{name}_train_log.csv"
        if not log.exists():
            return None
        return sum(float(l.split(",")[5]) for l in log.read_text().splitlines()[1:])
    txt = (EXP / "yolo_train.log").read_text(encoding="utf-8", errors="ignore")
    m = re.search(rf">>> {name} done in ([\d.]+) min", txt)
    return float(m.group(1)) if m else None


# ---------------- table data ----------------
HEADERS = ["Method", "Backbone", "Params (M) \u2193", "FLOPs (G) \u2193", "P \u2191",
           "R \u2191", "mAP@0.5 \u2191", "mAP@0.5:0.95 \u2191", "FPS \u2191"]
rows = []
for name in MODELS:
    disp, bb = DISPLAY[name]
    m, b = metrics[name], bench[name]
    rows.append({
        "name": name, "method": disp, "backbone": bb,
        "params": round(b["params_M"], 2),
        "flops": round(b["gflops"], 1),
        "p": round(m["precision"], 4),
        "r": round(m["recall"], 4),
        "map50": round(m["ap50"], 4),
        "map5095": round(m["ap50_95"], 4),
        "fps": round(b["fps"], 1),
        "train_min": train_time(name),
    })
COLS = ["params", "flops", "p", "r", "map50", "map5095", "fps"]
ASC = {"params": True, "flops": True, "p": False, "r": False,
       "map50": False, "map5095": False, "fps": False}  # True = smaller better

def fmt(v, col):
    if col in ("params", "flops", "fps"):
        return f"{v:.1f}" if col != "params" else f"{v:.2f}"
    return f"{v:.4f}"

best, second = {}, {}
for col in COLS:
    vals = sorted(rows, key=lambda r: r[col], reverse=not ASC[col])
    best[col] = vals[0]["name"]; second[col] = vals[1]["name"]

def cell_text(r, col):
    t = fmt(r[col], col)
    if r["name"] == best[col]:
        return t, "bold"
    if r["name"] == second[col]:
        return t + " \u2020", "normal"
    return t, "normal"

# ---------------- markdown / csv ----------------
md = ["| " + " | ".join(HEADERS) + " |",
      "|" + "---|" * len(HEADERS)]
for r in rows:
    cells = []
    for col in COLS:
        t = fmt(r[col], col)
        if r["name"] == best[col]:
            t = f"**{t}**"
        elif r["name"] == second[col]:
            t += " \u2020"
        cells.append(t)
    md.append("| " + f" | ".join([r["method"], r["backbone"]] + cells) + " |")
md.append("")
md.append("The best results are in **bold**, the second-best are marked with \u2020. "
          "All models were fine-tuned from COCO-pretrained weights (100 epochs for YOLO/RT-DETR; "
          "60 for SSD300; 50 for Faster R-CNN, batch sizes 4-16 on an RTX 3060 Laptop 6GB). "
          "FLOPs are computed at 640\u00d7640 input except SSD300 (300\u00d7300). "
          "FPS measured end-to-end (including NMS), batch size 1, FP16, on the same GPU. "
          "All metrics are evaluated on the same validation set (203 images) with the same COCO evaluator.")
(OUT_DIR / "comparison_table.md").write_text("\n".join(md) + "\n", encoding="utf-8")

import csv
with open(OUT_DIR / "comparison_table.csv", "w", newline="", encoding="utf-8-sig") as f:
    w = csv.writer(f)
    w.writerow(HEADERS)
    for r in rows:
        w.writerow([r["method"], r["backbone"]] + [fmt(r[c], c) for c in COLS])

# ---------------- xlsx ----------------
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "comparison"
ws.append(HEADERS)
for r in rows:
    ws.append([r["method"], r["backbone"]] + [fmt(r[c], c) for c in COLS])
for j, col in enumerate(COLS, start=3):
    for i, r in enumerate(rows, start=2):
        cell = ws.cell(row=i, column=j)
        if r["name"] == best[col]:
            cell.font = XLFont(bold=True)
        elif r["name"] == second[col]:
            cell.value += " \u2020"
ws2 = wb.create_sheet("training_details")
ws2.append(["Method", "Epochs", "Training time (min)", "Input size"])
for r in rows:
    ep = 100 if r["name"].startswith("yolo") or r["name"].startswith("rtdetr") \
        else (60 if r["name"] == "ssd" else 50)
    ws2.append([r["method"], ep, round(r["train_min"], 1) if r["train_min"] else None,
                300 if r["name"] == "ssd" else 640])
wb.save(OUT_DIR / "comparison_table.xlsx")

# ---------------- publication-style table PNG (booktabs) ----------------
ncols = len(HEADERS); nrows = len(rows) + 1
fig, ax = plt.subplots(figsize=(9.5, 2.6))
ax.axis("off")
for r in rows:
    pass
cw = [2.05, 1.85] + [1.05] * (ncols - 2)
xs = np.concatenate([[0], np.cumsum(cw)]) / sum(cw)
y_head, y_row = 0.88, 0.74
for j, h in enumerate(HEADERS):
    ax.text((xs[j] + xs[j + 1]) / 2, y_head, h, ha="center", va="center",
            fontsize=10.5, weight="bold")
for i, r in enumerate(rows):
    y = y_row - i * (0.74 / nrows)
    vals = [(r["method"], "normal"), (r["backbone"], "normal")] + [cell_text(r, c) for c in COLS]
    for j, (t, wt) in enumerate(vals):
        ax.text((xs[j] + xs[j + 1]) / 2, y, t, ha="center", va="center",
                fontsize=10, weight=wt)
for y in [y_head + 0.045, y_row - (nrows - 1) * (0.74 / nrows) - 0.035]:
    ax.plot([0, 1], [y, y], color=INK, lw=1.4)
ax.plot([0, 1], [y_row + 0.03, y_row + 0.03], color=INK, lw=0.8)
ax.text(0, -0.12, "The best results are in bold; \u2020 denotes the second-best. All metrics are evaluated on the "
        "same validation set (203 images) using the same COCO evaluator.", fontsize=7.5, color=MUTED, va="top")
ax.set_xlim(0, 1); ax.set_ylim(-0.15, 1.0)
fig.savefig(OUT_DIR / "table_comparison.png", bbox_inches="tight")
plt.close(fig)

# ---------------- metrics bar panels ----------------
fig, axes = plt.subplots(1, 3, figsize=(10.5, 3.2))
panels = [("mAP@0.5", "map50", 1.0), ("mAP@0.5:0.95", "map5095", 1.0), ("FPS", "fps", None)]
labels = [DISPLAY[n][0].replace("YOLO", "YOLO") for n in MODELS]
for ax, (title, col, ymax) in zip(axes, panels):
    vals = [r[col] for r in rows]
    bars = ax.bar(range(len(MODELS)), vals, color=[COLOR[n] for n in MODELS],
                  width=0.62, zorder=3)
    for b_, v in zip(bars, vals):
        ax.text(b_.get_x() + b_.get_width() / 2, v + (0.02 if col != "fps" else max(vals) * 0.02),
                f"{v:.1f}" if col == "fps" else f"{v:.4f}",
                ha="center", va="bottom", fontsize=8.5)
    ax.set_title(title, fontsize=11)
    ax.set_xticks(range(len(MODELS)))
    ax.set_xticklabels(labels, rotation=28, ha="right", fontsize=8.5)
    ax.set_ylim(0, ymax if ymax else max(vals) * 1.25)
    if col == "fps":
        ax.set_ylim(0, max(vals) * 1.3)
    ax.grid(axis="y", color=GRID, lw=0.8, zorder=0)
    for s in ["top", "right"]:
        ax.spines[s].set_visible(False)
fig.tight_layout()
fig.savefig(OUT_DIR / "metrics_bars.png", bbox_inches="tight")
plt.close(fig)

# ---------------- PR curves ----------------
fig, ax = plt.subplots(figsize=(5.2, 4.6))
for name in MODELS:
    d = np.load(EVAL_DIR / f"pr_{name}.npz")
    rec, prec = d["recall"], d["precision"]
    ap = metrics[name]["ap50"]
    ax.plot(rec, prec, color=COLOR[name], lw=2,
            label=f"{DISPLAY[name][0]} (AP={ap:.4f})")
    ax.annotate(DISPLAY[name][0], xy=(rec[-1], prec[-1]), xytext=(4, 0),
                textcoords="offset points", fontsize=7.5, color=INK, va="center")
ax.set_xlabel("Recall"); ax.set_ylabel("Precision")
ax.set_title("Precision-Recall curves (IoU = 0.5)", fontsize=11)
ax.legend(fontsize=7.5, loc="lower left", frameon=True, edgecolor=GRID)
ax.grid(color=GRID, lw=0.8)
for s in ["top", "right"]:
    ax.spines[s].set_visible(False)
fig.tight_layout()
fig.savefig(OUT_DIR / "pr_curves.png", bbox_inches="tight")
plt.close(fig)

print("ALL FIGURES/TABLES GENERATED", flush=True)
